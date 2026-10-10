"""Create an editable Word manuscript and native-shape specifications for Visio.

Run with the bundled Codex Python. Office automation is confined to the companion
PowerShell script; original Markdown, CSVs, checkpoints and exports are read only.
"""
from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'docs/PAPER_DEPTH_ESTIMATION_WITH_VISUALIZATIONS.md'
OUT = ROOT/'docs/paper_word_visio'
DOCX = OUT/'大视差与复杂遮挡深度估计论文.docx'
CONFIGS = [('range','Base'),('oa_range','Base+A'),('vis','Base+B'),
           ('range_hyp','Base+C'),('oa','Base+A+B'),('oa_full','Base+A+C'),
           ('hyp','Base+B+C'),('oa_hyp','Base+A+B+C')]


def xml(tag, **attrs):
    node = OxmlElement(tag)
    for k,v in attrs.items():
        node.set(qn(k), str(v))
    return node


SYMBOLS = dict(zip(
    'Delta Omega beta cap delta ell epsilon eta gamma ge geq in kappa lambda ldots le leq lfloor mu nu phi pi rfloor rho sigma star subseteq tau times top xi zeta'.split(),
    'Δ Ω β ∩ δ ℓ ε η γ ≥ ≥ ∈ κ λ … ≤ ≤ ⌊ μ ν φ π ⌋ ρ σ ⋆ ⊆ τ × ⊤ ξ ζ'.split()))
CALLIGRAPHIC = {'L':'ℒ','R':'ℛ','H':'ℋ','E':'ℰ','D':'𝒟','O':'𝒪'}


class MathParser:
    """The manuscript's LaTeX subset mapped to structured, editable Office Math."""
    def __init__(self, value):
        self.value = value
        self.pos = 0

    def run(self, text, roman=False, bold=False):
        r = xml('m:r')
        prop = xml('m:rPr')
        if roman:
            prop.append(xml('m:nor'))
        elif bold:
            prop.append(xml('m:sty', **{'m:val':'b'}))
        r.append(prop)
        t = xml('m:t')
        t.text = text
        r.append(t)
        return r

    def skip(self):
        while self.pos < len(self.value) and self.value[self.pos].isspace():
            self.pos += 1

    def nodes(self, stop=None):
        result = []
        while self.pos < len(self.value):
            self.skip()
            if self.pos >= len(self.value): break
            if stop and self.value[self.pos] == stop:
                self.pos += 1
                break
            if self.value.startswith(r'\right', self.pos):
                break
            atom = self.atom()
            if atom is not None:
                result.append(self.scripts(atom))
        return result

    def group(self):
        self.skip()
        if self.pos < len(self.value) and self.value[self.pos] == '{':
            self.pos += 1
            return self.nodes('}')
        return [self.atom()]

    def raw_group(self):
        self.skip()
        if self.pos >= len(self.value) or self.value[self.pos] != '{':
            return ''
        self.pos += 1
        start = self.pos; depth = 1
        while self.pos < len(self.value) and depth:
            c=self.value[self.pos]
            depth += (c == '{')-(c == '}')
            self.pos += 1
        return self.value[start:self.pos-1]

    def container(self, tag, children):
        e=xml(tag)
        e.extend(x for x in children if x is not None)
        return e

    def atom(self):
        self.skip()
        if self.pos >= len(self.value): return self.run('')
        c=self.value[self.pos]; self.pos += 1
        if c == '{':
            return self.container('m:e', self.nodes('}'))
        if c != '\\':
            return self.run(c)
        match=re.match(r'[A-Za-z]+', self.value[self.pos:])
        if not match:
            if self.pos >= len(self.value): return self.run('')
            escaped=self.value[self.pos]; self.pos += 1
            return self.run({'!':'', ',':'\u2009', ';':'\u2005', '%':'%', '|':'‖'}.get(escaped,escaped))
        command=match.group(); self.pos += len(command)
        if command in SYMBOLS: return self.run(SYMBOLS[command])
        if command in ('quad','qquad'): return self.run('\u2003'*(1 if command=='quad' else 2))
        if command in ('exp','log','tanh','max','min','arg'):
            return self.run(command, roman=True)
        if command == 'operatorname':
            return self.run(self.raw_group(), roman=True)
        if command in ('mathbf','mathrm','mathcal','mathbb'):
            content=self.group()
            for node in content:
                for t in node.iter(qn('m:t')):
                    if command=='mathcal': t.text=''.join(CALLIGRAPHIC.get(ch,ch) for ch in (t.text or ''))
                    if command=='mathbb': t.text=''.join({'R':'ℝ','1':'𝟙'}.get(ch,ch) for ch in (t.text or ''))
                for r in node.iter(qn('m:r')):
                    prop=r.find(qn('m:rPr'))
                    if command=='mathrm': prop.append(xml('m:nor'))
                    if command=='mathbf': prop.append(xml('m:sty', **{'m:val':'b'}))
            return self.container('m:e', content)
        if command == 'frac':
            n=self.group(); d=self.group(); f=xml('m:f')
            f.append(self.container('m:num',n)); f.append(self.container('m:den',d))
            return f
        if command == 'sqrt':
            r=xml('m:rad'); p=xml('m:radPr'); p.append(xml('m:degHide', **{'m:val':'1'})); r.append(p)
            r.append(xml('m:deg')); r.append(self.container('m:e',self.group())); return r
        if command in ('hat','bar','widetilde'):
            a=xml('m:acc'); p=xml('m:accPr'); p.append(xml('m:chr',**{'m:val':{'hat':'̂','bar':'̅','widetilde':'̃'}[command]}))
            a.append(p); a.append(self.container('m:e',self.group())); return a
        if command=='sum':
            sub,sup=self.read_scripts()
            n=xml('m:nary'); p=xml('m:naryPr')
            p.append(xml('m:chr',**{'m:val':'∑'})); p.append(xml('m:limLoc',**{'m:val':'undOvr'}))
            p.append(xml('m:subHide',**{'m:val':str(int(sub is None))})); p.append(xml('m:supHide',**{'m:val':str(int(sup is None))}))
            n.append(p); n.append(self.container('m:sub',sub or [])); n.append(self.container('m:sup',sup or []))
            n.append(self.container('m:e',[self.scripts(self.atom())])); return n
        if command=='left':
            begin=self.delimiter(); content=self.nodes()
            if not self.value.startswith(r'\right',self.pos):
                raise ValueError('Unmatched left delimiter: '+self.value)
            self.pos += len(r'\right'); end=self.delimiter()
            d=xml('m:d'); prop=xml('m:dPr')
            prop.append(xml('m:begChr',**{'m:val':begin})); prop.append(xml('m:endChr',**{'m:val':end}))
            d.append(prop); d.append(self.container('m:e',content)); return d
        raise ValueError('Unsupported LaTeX command '+command+' in '+self.value)

    def delimiter(self):
        self.skip()
        c=self.value[self.pos]; self.pos += 1
        if c=='\\':
            c=self.value[self.pos]; self.pos += 1
            c={'|':'‖'}.get(c,c)
        return '' if c=='.' else c

    def read_scripts(self):
        sub=sup=None
        while True:
            self.skip()
            if self.pos >= len(self.value) or self.value[self.pos] not in '_^': break
            kind=self.value[self.pos]; self.pos += 1
            value=self.group()
            if kind=='_': sub=value
            else: sup=value
        return sub,sup

    def scripts(self, atom):
        sub,sup=self.read_scripts()
        if sub is None and sup is None: return atom
        node=xml('m:sSubSup' if sub is not None and sup is not None else 'm:sSub' if sub is not None else 'm:sSup')
        node.append(self.container('m:e',[atom]))
        if sub is not None: node.append(self.container('m:sub',sub))
        if sup is not None: node.append(self.container('m:sup',sup))
        return node

    def parse(self):
        math=xml('m:oMath'); math.extend(self.nodes())
        if self.pos != len(self.value): raise ValueError('Unparsed equation tail '+self.value[self.pos:])
        # Grouping is an implementation detail, not an OMML element at top level.
        for node in list(math.iter(qn('m:e'))):
            parent=node.getparent()
            if parent.tag in (qn('m:oMath'),qn('m:e'),qn('m:num'),qn('m:den'),qn('m:sub'),qn('m:sup')):
                index=parent.index(node)
                children=list(node)
                parent.remove(node)
                for offset, child in enumerate(children): parent.insert(index+offset,child)
        return math


def parse_blocks(text):
    lines=text.splitlines(); i=0; result=[]
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line=='$$':
            j=i+1
            while j<len(lines) and lines[j].strip()!='$$': j+=1
            value='\n'.join(lines[i+1:j]); number=re.search(r'\\tag\{(\d+)\}',value).group(1)
            result.append(('equation',re.sub(r'\\tag\{\d+\}','',value).strip(),number)); i=j+1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells): rows.append(cells)
                i+=1
            result.append(('table',rows)); continue
        h=re.match(r'^(#{1,4})\s+(.*)$',line)
        if h: result.append(('heading',len(h.group(1)),h.group(2))); i+=1; continue
        image=re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)',line)
        if image: result.append(('image',image.group(1),image.group(2))); i+=1; continue
        if re.match(r'^\*\*(图|表)\d+\.',line):
            result.append(('caption',line)); i+=1; continue
        result.append(('paragraph',line)); i+=1
    return result


def fonts(style, size, east='宋体', bold=False):
    style.font.name='Times New Roman'; style.font.size=Pt(size)
    style.font.bold=bold; style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)


def inline(paragraph,text):
    pattern=r'(\$[^$]+\$|\*\*.*?\*\*|\[[^]]+\]\([^)]+\)|`[^`]+`)'
    for part in re.split(pattern,text):
        if not part: continue
        if part.startswith('$') and part.endswith('$'):
            paragraph._p.append(MathParser(part[1:-1]).parse())
        elif part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold=True
        elif re.fullmatch(r'\[[^]]+\]\([^)]+\)',part):
            m=re.fullmatch(r'\[([^]]+)\]\(([^)]+)\)',part)
            label,url=m.groups(); run=paragraph.add_run(label)
            if url.startswith('https://'):
                rid=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
                h=xml('w:hyperlink',**{'r:id':rid}); paragraph._p.remove(run._r); h.append(run._r); paragraph._p.append(h)
        else:
            paragraph.add_run(part[1:-1] if part.startswith('`') and part.endswith('`') else part)


def split_math(value):
    # Split only top-level expression separators; never split fraction arguments.
    pieces=[]; start=0; depth=0; i=0
    while i<len(value):
        if value[i]=='{' and (i==0 or value[i-1]!='\\'): depth+=1
        elif value[i]=='}' and (i==0 or value[i-1]!='\\'): depth-=1
        if depth==0 and (value.startswith(r'\qquad',i) or value.startswith(r'\quad',i)):
            token=r'\qquad' if value.startswith(r'\qquad',i) else r'\quad'
            if i-start>45:
                pieces.append(value[start:i].strip()); start=i+len(token)
            i+=len(token); continue
        i+=1
    pieces.append(value[start:].strip())
    return [p for p in pieces if p]


class PaperBuilder:
    def __init__(self):
        self.doc=Document(); self.landscape=False; self.appendix_a=False
        self.report={'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
                     'display_equations':0,'inline_equations':0,'tables':0,'figures':0,'source_blocks':0}
        styles=self.doc.styles
        fonts(styles['Normal'],12)
        styles['Normal'].paragraph_format.line_spacing=1.4
        styles['Normal'].paragraph_format.space_after=Pt(5)
        styles['Normal'].paragraph_format.first_line_indent=Pt(24)
        for level,size in [(1,15),(2,13),(3,12),(4,11.5)]:
            s=styles[f'Heading {level}']; fonts(s,size,'黑体',True)
            s.paragraph_format.keep_with_next=True
            s.paragraph_format.first_line_indent=Pt(0)
            s.paragraph_format.space_before=Pt(12 if level==1 else 9)
            s.paragraph_format.space_after=Pt(6)
        fonts(styles['Title'],20,'黑体',True)
        for border in list(styles['Title'].element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        styles['Title'].paragraph_format.first_line_indent=Pt(0)
        styles['Title'].paragraph_format.space_after=Pt(14)
        fonts(styles['Caption'],10,'宋体')
        styles['Caption'].paragraph_format.first_line_indent=Pt(0)
        styles['Caption'].paragraph_format.space_before=Pt(5)
        styles['Caption'].paragraph_format.space_after=Pt(8)
        styles['Caption'].paragraph_format.keep_with_next=False
        sec=self.doc.sections[0]; self.page_setup(sec,False)
        footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent=Pt(0)
        field=xml('w:fldSimple',**{'w:instr':'PAGE'}); field.append(xml('w:r')); footer._p.append(field)
        self.doc.core_properties.title='面向大视差与复杂遮挡的 Vis-MVSNet 深度估计改进方法'
        self.doc.core_properties.subject='参考视图深度估计与困难区域实验'
        self.doc.core_properties.author=''

    def page_setup(self,sec,land):
        sec.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
        sec.page_width=Inches(11 if land else 8.5); sec.page_height=Inches(8.5 if land else 11)
        sec.top_margin=sec.bottom_margin=Inches(.65)
        sec.left_margin=sec.right_margin=Inches(.7)
        sec.header_distance=sec.footer_distance=Inches(.3)

    def orient(self,land,force_page=False):
        if land!=self.landscape:
            sec=self.doc.add_section(WD_SECTION_START.NEW_PAGE); self.page_setup(sec,land)
            self.landscape=land
        elif force_page:
            p=self.doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
            p.add_run().add_break(WD_BREAK.PAGE)

    def picture(self,path,width,alt):
        p=self.doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(2)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        shape=p.add_run().add_picture(str(path),width=Inches(width))
        shape._inline.docPr.set('descr',alt)

    def grid_figure(self,path,alt):
        match=re.search(r'(View[35])/(?:scenes/)?(large_disparity_dominant|occlusion_dominant|joint_difficulty)/(scan[^/]+)/',path.as_posix())
        series,category,key=match.groups()
        panel_dir=ROOT/'outputs/paper_visualizations_ranked'/series/'scenes'/category/key/'panels'
        panels=[('参考图像','reference_rgb'),('GT深度','gt_depth'),('GT重投影位移','large_disparity_map'),
                ('源视图遮挡比例','source_occlusion_ratio'),('困难区域叠加','difficulty_overlay'),
                ('Base深度','base_depth'),('完整方法深度','full_model_depth'),('Base绝对误差','base_abs_error'),
                ('完整方法绝对误差','full_model_abs_error'),('红改善 蓝退化','error_gain'),
                ('至少一源遮挡','occluded_any_mask'),('至少两源遮挡','occluded_ge2_mask'),('大视差掩码','large_disparity_mask'),
                ('大视差与一源遮挡','large_disp_and_occluded_any_mask'),('大视差与两源遮挡','large_disp_and_occluded_ge2_mask')]
        self.orient(True,force_page=not getattr(self,'grid_header_ready',False))
        self.grid_header_ready=False
        table=self.doc.add_table(rows=3,cols=5); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        for i,(title,file) in enumerate(panels):
            cell=table.cell(i//5,i%5); cell.width=Inches(1.91)
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.keep_with_next=True; p.paragraph_format.line_spacing=1
            r=p.add_run(title); r.font.size=Pt(9); r.bold=True
            p=cell.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.line_spacing=1
            p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True
            shape=p.add_run().add_picture(str(panel_dir/(file+'.png')),width=Inches(1.82))
            shape._inline.docPr.set('descr',alt+' '+title)
        self.style_table(table,figure=True)
        legend=self.doc.add_paragraph(style='Caption')
        legend.paragraph_format.keep_with_next=True
        legend.add_run('RGB叠加：红色仅大视差  绿色仅至少两源遮挡  黄色交集；误差差值：红色改善  蓝色退化；黄色线为GT困难区域轮廓。').font.size=Pt(9)

    def style_table(self,table,figure=False):
        for row_idx,row in enumerate(table.rows):
            pr=row._tr.get_or_add_trPr(); pr.append(xml('w:cantSplit'))
            if row_idx==0 and not figure: pr.append(xml('w:tblHeader'))
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr(); borders=xml('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    borders.append(xml('w:'+edge,**{'w:val':'single','w:sz':'3','w:color':'D9D9D9'}))
                tcpr.append(borders)
                margins=xml('w:tcMar')
                for edge in ('top','bottom'): margins.append(xml('w:'+edge,**{'w:w':'60','w:type':'dxa'}))
                for edge in ('left','right'): margins.append(xml('w:'+edge,**{'w:w':'70','w:type':'dxa'}))
                tcpr.append(margins)
                if row_idx==0 and not figure:
                    tcpr.append(xml('w:shd',**{'w:fill':'E8EDF2','w:val':'clear'}))

    def data_table(self,rows):
        cols=len(rows[0])
        panel_order = rows[0] == ['行','第一列','第二列','第三列','第四列','第五列']
        if any(len(r)!=cols for r in rows): raise ValueError('Unequal table rows')
        if cols>=8: self.orient(True)
        available=9.6 if self.landscape else 7.1
        if panel_order: widths=[.55]+[(available-.55)/5]*5
        elif cols==11: widths=[1.6]+[.8]*10
        elif cols==10: widths=[.52,1.6,1.36,.55,.83,.83,.87,.82,.82,.80]
        elif cols==8: widths=[1.55]+[(available-1.55)/7]*7
        elif cols==2: widths=[available*.36,available*.64]
        else: widths=[available/cols]*cols
        scale=available/sum(widths); widths=[x*scale for x in widths]
        table=self.doc.add_table(rows=len(rows),cols=cols)
        table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        for i,row in enumerate(rows):
            for j,content in enumerate(row):
                cell=table.cell(i,j); cell.width=Inches(widths[j]); p=cell.paragraphs[0]
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (cols>2 or j==0) else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.space_before=Pt(1); p.paragraph_format.line_spacing=1 if panel_order else 1.12
                if 4<=len(rows)<=13:
                    p.paragraph_format.keep_with_next=i<len(rows)-1
                inline(p,content)
                for r in p.runs: r.font.size=Pt(9 if cols>=8 or panel_order else 10); r.bold=(i==0)
        self.style_table(table)
        spacer=self.doc.add_paragraph()
        spacer.paragraph_format.space_after=Pt(2)
        if panel_order:
            for row in table.rows:
                for cell in row.cells:
                    for item in cell._tc.xpath('./w:tcPr/w:tcMar/w:top | ./w:tcPr/w:tcMar/w:bottom'):
                        item.set(qn('w:w'),'20')
            spacer.paragraph_format.line_spacing=Pt(1)
            spacer.paragraph_format.space_after=Pt(0)
        self.report['tables']+=1

    def build(self):
        blocks=parse_blocks(SOURCE.read_text(encoding='utf8')); self.report['source_blocks']=len(blocks)
        for i,block in enumerate(blocks):
            kind=block[0]
            if kind=='heading':
                level,title=block[1:]
                if level==4 and i+1<len(blocks) and blocks[i+1][0]=='image' and blocks[i+1][2].endswith('overview_3x5.png'):
                    self.orient(True,force_page=True)
                    self.grid_header_ready=True
                if level==1:
                    p=self.doc.add_paragraph(title,style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                else:
                    if level==2:
                        self.appendix_a=title.startswith('附录 A')
                        self.in_conclusion=title.startswith('7 结论')
                        self.orient(self.appendix_a or title.startswith(('附录 B','5 实验结果')))
                    elif level==3 and not self.appendix_a:
                        self.orient(title.startswith(('5.1 ','5.2 ')))
                    self.doc.add_paragraph(title,style=f'Heading {min(level-1,4)}')
            elif kind=='equation':
                table=self.doc.add_table(rows=1,cols=2)
                table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
                width=9.6 if self.landscape else 7.1
                table.cell(0,0).width=Inches(width-.6); table.cell(0,1).width=Inches(.6)
                table.rows[0]._tr.get_or_add_trPr().append(xml('w:cantSplit'))
                pieces=split_math(block[1])
                for idx,part in enumerate(pieces):
                    cell=table.cell(0,0)
                    p=cell.add_paragraph() if idx else cell.paragraphs[0]
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.keep_with_next=True
                    mathpara=xml('m:oMathPara'); props=xml('m:oMathParaPr')
                    props.append(xml('m:jc',**{'m:val':'center'})); mathpara.append(props)
                    mathpara.append(MathParser(part).parse()); p._p.append(mathpara)
                num=table.cell(0,1); num.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=num.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.add_run('('+block[2]+')').font.size=Pt(10)
                for cell in table.rows[0].cells:
                    borders=xml('w:tcBorders')
                    for edge in ('top','left','bottom','right'):
                        borders.append(xml('w:'+edge,**{'w:val':'nil'}))
                    cell._tc.get_or_add_tcPr().append(borders)
                spacer=self.doc.add_paragraph()
                spacer.paragraph_format.space_after=Pt(0); spacer.paragraph_format.space_before=Pt(0)
                spacer.paragraph_format.line_spacing=Pt(1); spacer.add_run().font.size=Pt(1)
                self.report['display_equations']+=1
            elif kind=='table': self.data_table(block[1])
            elif kind=='image':
                alt,relative=block[1:]; path=SOURCE.parent/relative
                if path.name=='overview_3x5.png': self.grid_figure(path,alt)
                else:
                    is_diagram=path.name in ('framework.png','modules_abc_detailed.png')
                    if is_diagram:
                        self.orient(True,force_page=True)
                        page='01_总体框架' if path.name=='framework.png' else '02_ABC模块'
                        replacement=OUT/'visio_previews'/(page+'.png')
                        if not replacement.is_file(): raise FileNotFoundError(replacement)
                        path=replacement
                    self.picture(path,9.45 if self.landscape else 7.0,alt)
                self.report['figures']+=1
            elif kind=='caption':
                # Orient wide tables before their caption, keeping the pair together.
                if block[1].startswith('**表') and i+1<len(blocks) and blocks[i+1][0]=='table' and len(blocks[i+1][1][0])>=8:
                    self.orient(True)
                p=self.doc.add_paragraph(style='Caption'); inline(p,block[1])
                p.paragraph_format.keep_with_next=block[1].startswith('**表')
            else:
                p=self.doc.add_paragraph(); inline(p,block[1])
                if block[1].startswith('图7—图12对比 Base 与完整方法。'):
                    p.paragraph_format.keep_with_next=True
                if getattr(self,'in_conclusion',False) and i+1<len(blocks) and blocks[i+1][0]=='paragraph':
                    p.paragraph_format.keep_with_next=True
                if re.match(r'^\d+\.',block[1]): p.paragraph_format.first_line_indent=Pt(0)
                if re.match(r'^\[\d+\]',block[1]):
                    p.paragraph_format.first_line_indent=Pt(-18); p.paragraph_format.left_indent=Pt(18)
                    p.paragraph_format.line_spacing=1.15
                    for r in p.runs: r.font.size=Pt(10)
        self.doc.save(DOCX)
        root=self.doc._element
        self.report['native_math_objects']=len(root.findall('.//'+qn('m:oMath')))
        self.report['inline_equations']=len(re.findall(r'(?<!\$)\$(?!\$)(.+?)(?<!\$)\$(?!\$)',SOURCE.read_text(encoding='utf8')))
        assert self.report['display_equations']==26 and self.report['figures']==18
        self.report['editable_tables']=self.report['tables']
        self.report['embedded_picture_objects']=len(root.findall('.//'+qn('w:drawing')))
        assert self.report['embedded_picture_objects']==186
        (OUT/'word_content_validation.json').write_text(json.dumps(self.report,ensure_ascii=False,indent=2),encoding='utf8')
        print(json.dumps(self.report,ensure_ascii=False))


def page(name,width=16,height=9):
    return {'name':name,'width':width,'height':height,'shapes':[],'connectors':[]}


def box(p,id,x,y,w,h,text,fill='#eef2f6',font=16,stroke='#536477',bold=False):
    p['shapes'].append({'id':id,'x':x,'y':y,'w':w,'h':h,'text':text,
                        'fill':fill,'stroke':stroke,'font':font,'bold':bold})


def label(p,id,x,y,w,h,text,font=16,bold=False):
    box(p,id,x,y,w,h,text,'none',font,'none',bold)


def arrow(p,source,target,color='#536477',dashed=False,begin='right',end='left',head=True):
    p['connectors'].append({'from':source,'to':target,'color':color,'dashed':dashed,'begin':begin,'end':end,'head':head})


def routed_arrow(p,source,target,points,color,dashed=False,begin='right',end='left'):
    previous=source
    for index,(x,y) in enumerate(points):
        junction=f'route_{source}_{target}_{index}'
        box(p,junction,x-.005,y-.005,.01,.01,'','none',1,'none')
        arrow(p,previous,junction,color,dashed,begin=begin if index==0 else 'center',end='center',head=False)
        previous=junction
    arrow(p,previous,target,color,dashed,begin='center',end=end)


def diagram_spec():
    pages=[]
    p=page('01_总体框架',16,10); pages.append(p)
    label(p,'title',.4,9.5,15.2,.4,'面向大视差与复杂遮挡的三级深度估计框架',23,True)
    box(p,'input',.3,7.35,1.6,1.1,'参考与源图像\n相机参数',font=15)
    box(p,'features',2.2,7.25,2.1,1.3,'共享 2D U-Net\n多尺度特征\n1/8、1/4、1/2',font=15)
    for t,x in [(1,4.7),(2,8.1),(3,11.5)]:
        box(p,f's{t}',x,7.2,2.65,1.4,f'Stage {t}   '+['1/8','1/4','1/2'][t-1]+'\n逐源匹配与可靠性估计\nC 融合与阶段深度回归', '#eef5fb' if t==1 else '#fff3e5',14)
    box(p,'output',14.55,7.35,1.1,1.1,'最终\n深度',font=16)
    for a,b in [('input','features'),('features','s1'),('s1','s2'),('s2','s3'),('s3','output')]:arrow(p,a,b)
    box(p,'b2',7.8,8.85,2.6,.55,'B：前级均值与标准差 → 候选', '#fff0da',13,'#dc7900')
    box(p,'b3',11.25,8.85,2.6,.55,'B：前级均值与标准差 → 候选', '#fff0da',13,'#dc7900')
    arrow(p,'b2','s2','#dc7900',begin='bottom',end='top');arrow(p,'b3','s3','#dc7900',begin='bottom',end='top')
    box(p,'stage',.25,.25,15.5,6.55,'', '#ffffff',14)
    label(p,'inside_title',.55,6.1,14.5,.4,'单阶段内部结构  三个阶段采用相同的数据流',19,True)
    for id,x,y,w,h,text,fill in [
        ('candidates',.6,4.5,2.0,1.0,'参考与源特征\n候选深度 dₖᵗ','#e8f2fc'),
        ('warp',2.95,4.5,2.2,1.0,'单应性变换\n八组相关代价体 Fₛᵗ','#e8f2fc'),
        ('pair',5.5,4.5,2.25,1.0,'第一次 3D 正则化\n逐源隐变量 Vₛᵗ','#e7f6eb'),
        ('prob',8.1,4.5,2.25,1.0,'成对得分与概率 Pₛᵗ\n成对深度与熵 Eₛᵗ','#e7f6eb'),
        ('reliability',10.65,4.5,2.15,1.0,'共享熵编码\n对数尺度 uₛᵗ\n可见性 qₛᵗ','#e8f2fc'),
        ('fuse',8.0,2.9,3.0,1.05,'C：候选相关可靠性\nℓₛ,ₖᵗ = −uₛᵗ + rₛ,ₖᵗ\n源视图 softmax → wₛ,ₖᵗ','#d7f7ef'),
        ('regularize',11.5,2.9,3.2,1.05,'融合逐源隐变量\n第二次 3D 正则化\n深度概率 Pᵗ、均值 d̂ᵗ、标准差 σᵗ','#e7f6eb'),
        ('residual',4.3,2.9,3.1,1.05,'C 残差网络\n[Vₛᵗ, tanh aₛᵗ, Pₛᵗ, ξᵗ]\nrₛᵗ = ηᵗ tanh gC(Jₛᵗ)','#d7f7ef'),
        ('gt',.6,.95,2.3,1.15,'训练 GT 与相机\n重投影及深度一致性\n可见、遮挡、忽略标签','#e6efff'),
        ('a',3.4,.95,3.15,1.15,'A：逐源遮挡感知监督\n类别平衡 Focal BCE\n可见成对监督与误差停止梯度','#e6efff'),
        ('b',7.2,.95,3.6,1.15,'B：Stage 2/3 的候选生成\nhᵗ = clip(κᵗsᵗ, hminᵗ, hmaxᵗ)\n全局边界求交与均匀采样','#fff0da')]:
        box(p,id,x,y,w,h,text,fill,14)
    for a,b in [('candidates','warp'),('warp','pair'),('pair','prob'),('prob','reliability'),('residual','fuse'),('fuse','regularize'),('gt','a')]:arrow(p,a,b)
    arrow(p,'pair','residual','#008779',begin='bottom',end='top')
    arrow(p,'reliability','fuse','#008779',begin='bottom',end='top')
    routed_arrow(p,'a','reliability',[(4.975,2.4),(7.9,2.4),(7.9,5.85),(11.725,5.85)],'#2864da',True,begin='top',end='top')
    arrow(p,'regularize','b','#dc7900',begin='bottom',end='top')
    routed_arrow(p,'b','candidates',[(6.8,1.525),(6.8,.6),(.45,.6),(.45,4.15),(1.6,4.15)],'#dc7900',True,begin='left',end='bottom')
    label(p,'legend',11.3,.8,3.6,1.1,'蓝色：A 的训练监督\n橙色：B 的候选搜索\n青色：C 的候选相关融合',14)

    p=page('02_ABC模块',16,10); pages.append(p)
    label(p,'title',.3,9.3,15.4,.5,'A B C 模块的内部计算',23,True)
    modules=[
        ('A',.35,'逐源遮挡感知监督','#e6efff','#2864da',[
            '参考 GT、源 GT 与相机参数',
            '参考真值点重投影到源视图\n源投影深度 Zₛ 与源深度 G̃ₛ',
            '几何有效性与深度容差\nτₛ = max(τa, τr G̃ₛ)',
            '可见：|Zₛ−G̃ₛ| ≤ τₛ\n遮挡：Zₛ > G̃ₛ+τₛ\n其他位置忽略',
            '成对概率熵 Eₛᵗ → 共享编码 Hₛᵗ\n可见性 logit → sigmoid qₛᵗ',
            '有效标签上的类别平衡 Focal BCE\n联合成对深度与不确定性监督']),
        ('B',5.65,'自适应深度搜索范围','#fff0da','#dc7900',[
            '前级融合概率 Pᵗ⁻¹、候选深度',
            '概率期望与标准差\n前级结果双线性对齐为 μᵗ、sᵗ',
            '逐像素搜索半宽\nhᵗ = clip(κᵗsᵗ, hminᵗ, hmaxᵗ)',
            '与全局深度边界求交\nd₋ᵗ=max(dmin, μᵗ−hᵗ)\nd₊ᵗ=min(dmax, μᵗ+hᵗ)',
            '均匀生成 Nt 个候选\ndₖᵗ=d₋ᵗ+k(d₊ᵗ−d₋ᵗ)/(Nt−1)',
            '共同评价 coverage、width 与误差\n候选覆盖和采样间距存在取舍']),
        ('C',10.95,'深度假设感知源视图融合','#d7f7ef','#008779',[
            '逐源隐变量 Vₛᵗ 与基础 logit −uₛᵗ',
            '十一通道输入\n[Vₛᵗ, tanh aₛᵗ, Pₛᵗ, ξᵗ]',
            '3D Conv → GroupNorm → ReLU\n1×1×1 Conv → tanh\nrₛ,ₖᵗ = ηᵗ tanh ζₛ,ₖᵗ',
            '基础分数与残差相加\nℓₛ,ₖᵗ = −uₛᵗ + rₛ,ₖᵗ',
            '在源视图维度归一化\nwₛ,ₖᵗ = softmaxₛ(ℓₛ,ₖᵗ)',
            '逐候选加权融合 Vᵗ = Σₛ wₛ,ₖᵗ Vₛᵗ\n第二次正则化与阶段深度回归'])]
    for name,x,title,fill,color,steps in modules:
        label(p,name+'_title',x,8.6,4.7,.5,name+'  '+title,18,True)
        for i,text in enumerate(steps):
            id=f'{name}_{i}'; box(p,id,x,7.35-i*1.18,4.7,1.0,text,fill,14,color)
            if i: arrow(p,f'{name}_{i-1}',id,color,begin='bottom',end='top')
    for name,x,title,fill,color,steps in modules:
        q=page({'A':'03_A遮挡监督','B':'04_B自适应范围','C':'05_C假设融合'}[name],10,10); pages.append(q)
        label(q,'title',.5,9.0,9,.6,name+'  '+title,25,True)
        for i,text in enumerate(steps):
            id=f'{name}_{i}'; box(q,id,1.3,7.55-i*1.22,7.4,1.10,text,fill,18,color)
            if i: arrow(q,f'{name}_{i-1}',id,color,begin='bottom',end='top')
    p=page('06_训练推理流程',14,9); pages.append(p)
    label(p,'title',.4,8.2,13.2,.5,'训练与推理的数据流',24,True)
    for id,x,y,w,h,text,fill in [
        ('images',.5,6.3,2.4,1.0,'图像与相机\n训练及推理输入','#e8f2fc'),
        ('forward',3.5,6.1,4.2,1.4,'三级前向\n候选生成 → 逐源匹配与可靠性\n候选相关融合 → 阶段深度','#e7f6eb'),
        ('depth',8.3,6.3,2.3,1.0,'最终深度\n推理输出','#e8f2fc'),
        ('gt',.5,3.8,2.4,1.0,'参考与源 GT\n仅训练和诊断','#e6efff'),
        ('labels',3.5,3.6,4.2,1.4,'几何标签\n可见性监督与有效像素掩码\n候选最近 GT 处的可选监督','#e6efff'),
        ('loss',8.3,3.6,4.5,1.4,'联合训练目标\n融合 L1 + 成对 L1 + 不确定性 NLL\nA 的 Focal + 可选 C 监督','#fff0da'),
        ('update',8.3,1.2,4.5,1.2,'反向传播与参数更新\nGT 不作为推理融合输入','#fff0da')]:box(p,id,x,y,w,h,text,fill,18)
    for a,b in [('images','forward'),('forward','depth'),('gt','labels'),('labels','loss')]:arrow(p,a,b)
    arrow(p,'forward','loss',begin='bottom',end='top');arrow(p,'loss','update',begin='bottom',end='top')
    routed_arrow(p,'update','forward',[(7.9,1.8),(7.9,5.65),(5.6,5.65)],'#dc7900',True,begin='left',end='bottom')

    regions=[('full','整体'),('boundary','深度边界'),('large_disparity','大视差'),
             ('occluded_any','任一源遮挡'),('occluded_majority','至少半数源遮挡'),
             ('large_disp_and_occluded','大视差∩遮挡'),('boundary_and_occluded','边界∩遮挡')]
    data={}
    for series,folder in [('View5','ablation_test_light3'),('View3','ablation_test_view3_light3')]:
        data[series]={}
        for raw,label_text in CONFIGS:
            path=ROOT/'eval'/folder/f'{raw}_{series.lower()}'/'summary_metrics.csv'
            with path.open(encoding='utf-8-sig',newline='') as handle:
                data[series][label_text]={r['region']:r for r in csv.DictReader(handle) if r['aggregation']=='pixel_weighted'}
        q=page('07_View5区域改善' if series=='View5' else '08_View3区域改善',16,10);pages.append(q)
        label(q,'title',.4,9.2,15.2,.5,series+'  八组配置相对 Base 的区域 Abs 降幅',23,True)
        for j,(_,title) in enumerate(regions):label(q,f'h{j}',3.6+j*1.7,8.4,1.65,.55,title,15,True)
        for i,(_,label_text) in enumerate(CONFIGS):
            y=7.45-i*.86;label(q,f'r{i}',.5,y,3.0,.65,label_text,17,True)
            for j,(region,_) in enumerate(regions):
                b=float(data[series]['Base'][region]['abs']);v=float(data[series][label_text][region]['abs']);gain=100*(b-v)/b
                intensity=min(1,abs(gain)/35)
                rgb=(int(246-100*intensity),int(248-30*intensity),int(246-100*intensity)) if gain>=0 else (255,225,225)
                color='#'+''.join(f'{v:02x}' for v in rgb)
                box(q,f'c{i}_{j}',3.6+j*1.7,y,1.65,.65,f'{gain:+.2f}%',color,17,'#d9d9d9')
        label(q,'units',.6,.35,14.8,.6,'正值为 Abs 降低；负值为增加。每个单元格均为独立可编辑形状，数值来自原始像素加权评测。',15)
    q=page('09_两套消融汇总',16,10); pages.append(q)
    label(q,'title',.4,9.2,15.2,.5,'完整配置与 Base 的困难区域 Abs 对照',23,True)
    target=['full','large_disparity','occluded_any','large_disp_and_occluded']
    titles={'full':'整体','large_disparity':'大视差','occluded_any':'任一源遮挡','large_disp_and_occluded':'大视差∩遮挡'}
    for sidx,series in enumerate(('View5','View3')):
        x0=.7+sidx*7.8;label(q,f's{sidx}',x0,8.4,7,.5,series,21,True)
        for i,region in enumerate(target):
            y=6.7-i*1.65;b=float(data[series]['Base'][region]['abs']);f=float(data[series]['Base+A+B+C'][region]['abs'])
            label(q,f'l{sidx}_{i}',x0,y+.3,1.6,.65,titles[region],15,True)
            for j,(v,color,model) in enumerate([(b,'#aab9ca','Base'),(f,'#489888','完整')]):
                yy=y+.6-j*.58;box(q,f'bar{sidx}_{i}_{j}',x0+1.8,yy,4.0*v/65,.4,'',color,12,'none')
                label(q,f'v{sidx}_{i}_{j}',x0+1.8+4.0*v/65+.07,yy,1.6,.4,f'{model} {v:.3f}',14)
    label(q,'unit',.7,.4,14.6,.5,'Abs 单位 mm；所有条形、数值与标签均可在 Visio 中单独修改。',15)
    q=page('10_阈值准确率',16,10); pages.append(q)
    label(q,'title',.4,9.2,15.2,.5,'八组配置的整体阈值准确率',23,True)
    for sidx,series in enumerate(('View5','View3')):
        x0=.3+sidx*8;label(q,f's{sidx}',x0,8.5,7.6,.4,series,21,True)
        for j,name in enumerate(('Acc2','Acc4','Acc8')):label(q,f'h{sidx}_{j}',x0+3+j*1.5,7.8,1.45,.5,name,17,True)
        for i,(_,label_text) in enumerate(CONFIGS):
            y=6.85-i*.78;label(q,f'r{sidx}_{i}',x0,y,2.9,.6,label_text,16,True)
            for j,metric in enumerate(('acc2','acc4','acc8')):
                value=float(data[series][label_text]['full'][metric])*100
                box(q,f'c{sidx}_{i}_{j}',x0+3+j*1.5,y,1.45,.6,f'{value:.2f}%','#edf4f8',16,'#d9d9d9')
    label(q,'unit',.4,.4,15.2,.5,'Acc2/4/8 分别为误差严格小于 2/4/8 mm 的有效像素比例。',15)
    return {'pages':pages,'metadata':{'source':SOURCE.relative_to(ROOT).as_posix(),'native_shapes':True,
                                    'bitmap_predictions_not_recreated':True}}


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('command',choices=['spec','word','check'])
    args=parser.parse_args(); OUT.mkdir(parents=True,exist_ok=True)
    if args.command=='spec':
        (OUT/'visio_diagrams_spec.json').write_text(json.dumps(diagram_spec(),ensure_ascii=False,indent=2),encoding='utf8')
        print('Prepared ten native Visio pages')
    elif args.command=='word': PaperBuilder().build()
    else:
        import zipfile
        with zipfile.ZipFile(DOCX) as z:
            root=etree.fromstring(z.read('word/document.xml'))
            text=''.join(root.itertext())
            if re.search(r'\\(?:frac|mathbb|mathbf|operatorname|tag)\b',text): raise ValueError('Raw LaTeX remains')
            if len(root.findall('.//'+qn('m:oMath')))<135: raise ValueError('Missing native math')
        print('DOCX native math and package checks passed')


if __name__=='__main__': main()
